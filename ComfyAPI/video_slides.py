"""
影片轉投影片(.pptx)

管線: 字幕(既有/ASR) + 場景擷取圖 -> LLM 歸納成數頁大綱 -> python-pptx 輸出。
模型資訊(base_url/model/api_key)由前端帶入, 本模組不寫死。
"""
import json
import re
from pathlib import Path

import requests

import video_scene
import video_transcribe
import youtube_clip

_BT = chr(96)


def _find_subtitle(src, base):
    src = Path(src)
    cand = src.with_suffix(".srt")
    if cand.is_file():
        return cand
    cand = base / (src.stem + "_asr.srt")
    if cand.is_file():
        return cand
    return None


def _find_scene_images(src, base, limit):
    stem_slug = youtube_clip.slug(Path(src).stem, 24)
    prefix = "scenes_" + stem_slug + "_"
    cands = [d for d in base.iterdir() if d.is_dir() and d.name.startswith(prefix)]
    if not cands:
        return []
    cands.sort(key=lambda d: d.stat().st_mtime, reverse=True)
    scene_dir = cands[0]
    imgs = []
    for f in sorted(scene_dir.glob("s*_start.jpg")):
        m = re.match(r"s(\d+)_start", f.name)
        idx = int(m.group(1)) if m else len(imgs) + 1
        imgs.append({"id": idx, "path": str(f), "caption": ""})
        if len(imgs) >= limit:
            break
    return imgs


def _make_scene_images(src, base, ffmpeg_location, segments, log, progress_cb, cancel_event):
    import time as _time
    import uuid as _uuid
    scenes = video_scene.detect_scenes(str(src))
    if not scenes:
        return []
    ts = _time.strftime("%Y%m%d_%H%M%S")
    slug = youtube_clip.slug(Path(src).stem, 24)
    out_dir = base / ("scenes_" + slug + "_" + ts + "_" + _uuid.uuid4().hex[:4])
    out_dir.mkdir(parents=True, exist_ok=True)
    results = video_scene.extract_scene_frames(
        str(src), scenes, str(out_dir), ffmpeg_location=ffmpeg_location,
        log=log, progress_cb=progress_cb, cancel_event=cancel_event,
        segments=segments,
    )
    imgs = []
    for r in results:
        imgs.append({
            "id": r["index"],
            "path": str(out_dir / r["start_file"]),
            "caption": r.get("caption", "") or "",
        })
    return imgs


def _llm_chat(base_url, model, api_key, messages, temperature=0.4, max_tokens=2048,
             enable_thinking=False):
    url = base_url.rstrip("/") + "/chat/completions"
    headers = {"Authorization": "Bearer " + (api_key or ""), "Content-Type": "application/json"}
    payload = {"model": model, "messages": messages, "temperature": temperature,
               "max_tokens": max_tokens, "stream": False}
    payload["chat_template_kwargs"] = {"enable_thinking": bool(enable_thinking)}
    r = requests.post(url, headers=headers, json=payload, timeout=180)
    if r.status_code >= 400:
        raise RuntimeError("模型服務 HTTP " + str(r.status_code) + ": " + r.text[:300])
    data = r.json()
    choices = data.get("choices") or []
    if not choices:
        raise RuntimeError("模型無回應")
    text = (choices[0].get("message") or {}).get("content")
    if isinstance(text, list):
        text = "".join(p.get("text", "") for p in text if isinstance(p, dict))
    return (text or "").strip()


def _parse_slides_json(raw, max_slides):
    s = raw.strip()
    s = re.sub("^" + _BT * 3 + r"[a-zA-Z]*\n?", "", s)
    s = re.sub(r"\n?" + _BT * 3 + "$", "", s)
    i = s.find("{")
    j = s.rfind("}")
    if i == -1 or j == -1 or j <= i:
        return None
    try:
        obj = json.loads(s[i:j + 1])
    except json.JSONDecodeError:
        return None
    slides = obj.get("slides")
    if not isinstance(slides, list) or not slides:
        return None
    out = []
    for sl in slides[:max_slides]:
        if not isinstance(sl, dict):
            continue
        title = str(sl.get("title", "")).strip()
        bullets = sl.get("bullets")
        if not isinstance(bullets, list):
            bullets = [str(bullets)] if bullets else []
        bullets = [str(b).strip() for b in bullets if str(b).strip()]
        img = sl.get("image_id")
        img = int(img) if isinstance(img, (int, float)) else None
        out.append({"title": title, "bullets": bullets, "image_id": img})
    return out or None


def _outline(title, segments, images, base_url, model, api_key, max_slides, target=None, enable_thinking=False):
    seg_lines = []
    for idx, seg in enumerate(segments, 1):
        seg_lines.append("[" + str(idx) + "] " + str(int(seg["start"])) + "-" + str(int(seg["end"])) + "s " + seg["text"])
    img_lines = []
    for im in images:
        cap = im.get("caption") or ""
        img_lines.append("(id=" + str(im["id"]) + ") " + cap)
    target_line = ("本次簡報希望貼近「" + target + "」的風格與用語。\n" if target else "")
    sys = "你是簡報摘要助手。只輸出嚴格 JSON, 不要任何額外文字或說明。"
    user = (
        "影片標題: " + title + "\n\n" + (target_line) +
        "字幕段落(index 時間 文字):\n" + "\n".join(seg_lines) + "\n\n"
        "可用圖片(id 說明):\n" + "\n".join(img_lines) + "\n\n"
        "請歸納成 4 到 " + str(max_slides) + " 頁繁體中文簡報大綱。每頁含 title(短標題)、"
        "bullets(2-4 條重點, 條列)、image_id(由可用圖片中挑選, 無合適填 null)。\n"
        "嚴格只輸出如下 JSON, 不要 code fence:\n"
        '{"slides":[{"title":"...","bullets":["...","..."],"image_id":1}]}'
    )
    raw = _llm_chat(base_url, model, api_key,
                    [{"role": "system", "content": sys}, {"role": "user", "content": user}],
                    temperature=0.4, max_tokens=2048, enable_thinking=bool(enable_thinking))
    return _parse_slides_json(raw, max_slides)


def _fallback_slides(segments, images, max_slides):
    out = []
    for idx, seg in enumerate(segments[:max_slides]):
        img = images[idx]["id"] if images and idx < len(images) else None
        out.append({"title": "段落 " + str(idx + 1), "bullets": [seg["text"]], "image_id": img})
    return out


def _describe_image(image_path, base_url, model, api_key):
    import run_image2prompt as rip
    try:
        text = rip.generate(image_path=image_path, target=None, base_url=base_url,
                            model=model, api_key=api_key, enable_thinking=False,
                            log=lambda *_: None)
    except Exception:
        return ""
    text = (text or "").strip()
    text = re.sub(r"<think.*?</think>", "", text, flags=re.S).strip()
    first = text.splitlines()[0].strip() if text else ""
    return first[:200]


def _prepare(video_path, base, ffmpeg_location, base_url, model, api_key,
             language=None, max_slides=10, do_vision=False, target=None, enable_thinking=False,
             log=None, progress_cb=None, cancel_event=None):
    """字幕(既有/ASR) + 場景圖 + LLM 大綱。回傳 (title, slides, images)。"""
    def _log(*a):
        if log:
            log(*a)

    def _cancel():
        if cancel_event is not None and cancel_event.is_set():
            raise KeyboardInterrupt

    def _prog(frac):
        if progress_cb:
            progress_cb(int(frac * 100), 100)

    video_path = Path(video_path)
    base = Path(base)
    title = video_path.stem

    _log("取得字幕…")
    srt = _find_subtitle(video_path, base)
    if srt is None:
        _log("無字幕檔, 執行語音辨識…")
        out_srt = base / (video_path.stem + "_asr.srt")
        video_transcribe.transcribe(
            str(video_path), str(out_srt), ffmpeg_location,
            language=language, log=log,
            progress_cb=lambda v, m: _prog(0.05 + 0.35 * (v / max(1, m))),
            cancel_event=cancel_event)
        srt = out_srt
    _cancel()
    segments = video_scene.load_srt_segments(str(srt))
    if not segments:
        raise RuntimeError("字幕內容為空, 無法產生內容")

    _log("取得場景擷取圖…")
    _prog(0.42)
    images = _find_scene_images(video_path, base, max_slides)
    if not images:
        _log("無擷取圖, 偵測場景並抽圖…")
        images = _make_scene_images(
            video_path, base, ffmpeg_location, segments, log,
            progress_cb=lambda v, m: _prog(0.45 + 0.20 * (v / max(1, m))),
            cancel_event=cancel_event)
    _cancel()

    if do_vision and images:
        _log("辨識圖片內容…")
        for k, im in enumerate(images):
            _cancel()
            cap = _describe_image(im["path"], base_url, model, api_key)
            if cap:
                im["caption"] = cap
            _prog(0.65 + 0.10 * ((k + 1) / len(images)))

    _log("產生大綱…")
    _prog(0.78)
    slides = None
    try:
        slides = _outline(title, segments, images, base_url, model, api_key, max_slides,
                          target=target, enable_thinking=enable_thinking)
    except Exception as e:
        _log("大綱產生失敗, 改用字幕分段: " + str(e))
    if not slides:
        slides = _fallback_slides(segments, images, max_slides)
    _cancel()
    return title, slides, images


def build_slides(video_path, out_pptx, base, ffmpeg_location, base_url, model, api_key,
                 language=None, max_slides=10, do_vision=False, target=None, enable_thinking=False,
                 log=None, progress_cb=None, cancel_event=None):
    def _log(*a):
        if log:
            log(*a)

    def _prog(frac):
        if progress_cb:
            progress_cb(int(frac * 100), 100)

    title, slides, images = _prepare(
        video_path, base, ffmpeg_location, base_url, model, api_key,
        language=language, max_slides=max_slides, do_vision=do_vision,
        target=target, enable_thinking=enable_thinking,
        log=log, progress_cb=progress_cb, cancel_event=cancel_event)

    _log("寫入 .pptx…")
    _prog(0.92)
    _write_pptx(title, slides, images, Path(out_pptx))
    _prog(1.0)
    _log("完成: " + Path(out_pptx).name)
    return Path(out_pptx)


def _write_pptx(title, slides, images, out_pptx):
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from PIL import Image

    prs = Presentation()
    slide_w, slide_h = prs.slide_width, prs.slide_height

    tslide = prs.slides.add_slide(prs.slide_layouts[0])
    tslide.shapes.title.text = title
    try:
        tslide.placeholders[1].text = "共 " + str(len(slides)) + " 頁 · 由字幕摘要產生"
    except Exception:
        pass

    img_by_id = {im["id"]: im for im in images}

    for sl in slides:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(0.4), Inches(0.25), slide_w - Inches(0.8), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = sl.get("title") or ""
        p.font.size = Pt(28)
        p.font.bold = True

        img = img_by_id.get(sl.get("image_id")) if sl.get("image_id") is not None else None
        text_left = Inches(0.5)
        text_w = slide_w - Inches(1.0)
        if img and Path(img["path"]).is_file():
            pic_w = Inches(5.2)
            pic_h = Inches(3.0)
            top = Inches(1.2)
            left = Inches(0.5)
            try:
                with Image.open(img["path"]) as im0:
                    iw, ih = im0.size
            except Exception:
                iw, ih = 16, 9
            ratio = iw / ih if ih else 1.0
            w = pic_w
            h = Emu(int(w / ratio))
            if h > pic_h:
                h = pic_h
                w = Emu(int(h * ratio))
            s.shapes.add_picture(img["path"], left, top, width=w, height=h)
            text_left = left + w + Inches(0.4)
            text_w = slide_w - text_left - Inches(0.5)

        tb2 = s.shapes.add_textbox(text_left, Inches(1.2), text_w, slide_h - Inches(1.6))
        tf2 = tb2.text_frame
        tf2.word_wrap = True
        bullets = sl.get("bullets") or []
        for bi, b in enumerate(bullets):
            para = tf2.paragraphs[0] if bi == 0 else tf2.add_paragraph()
            para.text = b
            para.font.size = Pt(18)
            para.space_after = Pt(8)

    out_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_pptx))
